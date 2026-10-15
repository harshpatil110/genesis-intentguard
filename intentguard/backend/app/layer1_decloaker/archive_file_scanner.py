"""
IntentGuard Layer 1 Disguised Executables, Archives, & Poisoned Spreadsheets Scanner
====================================================================================
GENESIS 2.0 | Track: EVA 03 | PS 01
Separating Data from Authority: Securing AI Agents Against Indirect Prompt Injection

Deep-packet inspection of file headers (magic bytes), archive structures, and
spreadsheets to neutralize:
- Risk #1:  Malicious Attachments (Disguised PE/ELF executables, malicious formulas)
- Risk #7:  Poisoned Data Sources (Hidden Excel worksheets, DDE/formula injection)
- Risk #8:  Archive & File-Type Tricks (Double extensions, MIME mismatch, nested archives)
- Risk #19: Resource Exhaustion (Zip decompression bombs, massive uncompressed payloads)
"""

from __future__ import annotations

import base64
import hashlib
import io
import logging
import re
import zipfile
from typing import Any, Dict, List, Optional, Tuple, Union

import magic
import openpyxl

from app.layer1_decloaker.pdf_docx_decloaker import CloakedSpanFinding, DocumentDeCloakResult

logger = logging.getLogger(__name__)

# Max uncompressed archive limit: 25 MB
MAX_UNCOMPRESSED_ARCHIVE_BYTES = 25 * 1024 * 1024  # 25 MB

# Max compression ratio allowed before flagging as zip decompression bomb
MAX_ZIP_BOMB_RATIO = 50.0

# Nested archive extensions to inspect inside archives
_ARCHIVE_EXTENSIONS = frozenset({".zip", ".tar", ".gz", ".tgz", ".bz2", ".7z", ".rar", ".xz"})

# Executable MIME types and extensions
_EXECUTABLE_MIMES = frozenset(
    {
        "application/x-dosexec",
        "application/x-msdownload",
        "application/x-executable",
        "application/x-sharedlib",
        "application/x-pie-executable",
        "application/x-mach-binary",
        "application/x-elf",
    }
)

_EXECUTABLE_EXTENSIONS = frozenset(
    {
        "exe",
        "scr",
        "bat",
        "cmd",
        "com",
        "vbs",
        "vbe",
        "js",
        "jse",
        "wsf",
        "wsh",
        "ps1",
        "sh",
        "bin",
        "msi",
        "dll",
    }
)

_DOCUMENT_EXTENSIONS = frozenset(
    {
        "pdf",
        "docx",
        "doc",
        "xlsx",
        "xls",
        "csv",
        "png",
        "jpg",
        "jpeg",
        "webp",
        "gif",
        "txt",
        "json",
        "xml",
        "html",
        "htm",
        "mp3",
        "mp4",
        "wav",
    }
)

# Malicious spreadsheet formula patterns
_MALICIOUS_FORMULA_RE = re.compile(
    r"^[\s=\+\-@](?:WEBSERVICE|HYPERLINK|CMD|EXEC|SHELL|DDE|SYSTEM)\b",
    re.IGNORECASE,
)


class ArchiveFileScanner:
    """
    Inspection engine for:
    1. Magic bytes & executable disguises (.pdf.exe, MZ/ELF header mismatches)
    2. Archive safety (Zip decompression bombs, size limits, nested archives)
    3. Spreadsheet safety (Hidden worksheets, malicious DDE/Webservice formulas)
    """

    def scan_magic_bytes(
        self,
        file_bytes: bytes,
        filename: str,
    ) -> List[CloakedSpanFinding]:
        """
        Inspect the true file header magic bytes and detect disguised executables.
        """
        findings: List[CloakedSpanFinding] = []
        if not file_bytes:
            return findings

        # 1. Direct header magic byte inspection
        is_mz = file_bytes.startswith(b"MZ")
        is_elf = file_bytes.startswith(b"\x7fELF")
        is_macho = file_bytes[:4] in (
            b"\xfe\xed\xfa\xce",
            b"\xfe\xed\xfa\xcf",
            b"\xce\xfa\xed\xfe",
            b"\xcf\xfa\xed\xfe",
        )

        # 2. Extract MIME type via python-magic buffer inspection
        mime_type = ""
        try:
            mime_type = magic.from_buffer(file_bytes[:2048], mime=True) or ""
        except Exception as exc:
            logger.debug("magic.from_buffer notice: %s", exc)

        is_exec_mime = mime_type.lower() in _EXECUTABLE_MIMES

        # 3. Check for double extension tricks (e.g. invoice.pdf.exe)
        fn_clean = filename.lower().strip()
        parts = fn_clean.split(".")
        is_double_ext = False
        if len(parts) > 2:
            if parts[-1] in _EXECUTABLE_EXTENSIONS and parts[-2] in _DOCUMENT_EXTENSIONS:
                is_double_ext = True

        # 4. Disguise condition:
        # Document/Data extension on filename, but actual binary is PE/ELF/Mach-O or executable MIME
        is_doc_filename = any(fn_clean.endswith(f".{ext}") for ext in _DOCUMENT_EXTENSIONS)
        is_binary_executable = is_mz or is_elf or is_macho or is_exec_mime or is_double_ext

        if (is_doc_filename and is_binary_executable) or is_double_ext:
            detail = (
                f"Disguised executable detected: filename='{filename}', "
                f"mime='{mime_type}', MZ={is_mz}, ELF={is_elf}, double_ext={is_double_ext}"
            )
            findings.append(
                CloakedSpanFinding(
                    modality="file",
                    page_or_paragraph=1,
                    cloak_reason="DISGUISED_EXECUTABLE",
                    hidden_text=detail,
                    metadata_key="magic_bytes",
                    risk_category_ids=[1, 2, 8],
                )
            )

        return findings

    def scan_archive_safety(
        self,
        file_bytes: bytes,
        filename: str,
    ) -> Tuple[List[CloakedSpanFinding], str]:
        """
        Analyze archive compression ratio and member list without extracting to disk.
        """
        findings: List[CloakedSpanFinding] = []
        summary_text = ""

        is_zip = file_bytes.startswith(b"PK\x03\x04") or filename.lower().endswith(".zip")
        if not is_zip:
            return findings, summary_text

        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
                infolist = zf.infolist()
                total_uncompressed = sum(info.file_size for info in infolist)
                compressed_len = max(len(file_bytes), 1)
                ratio = float(total_uncompressed) / float(compressed_len)

                # Rule 1: Total uncompressed size limit (25 MB)
                if total_uncompressed > MAX_UNCOMPRESSED_ARCHIVE_BYTES:
                    findings.append(
                        CloakedSpanFinding(
                            modality="archive",
                            page_or_paragraph=1,
                            cloak_reason="ARCHIVE_EXCEEDS_SIZE_LIMIT",
                            hidden_text=(
                                f"Uncompressed archive size {total_uncompressed} bytes "
                                f"exceeds safety cap {MAX_UNCOMPRESSED_ARCHIVE_BYTES} bytes (25MB)"
                            ),
                            metadata_key="archive_size",
                            risk_category_ids=[8, 19],
                        )
                    )

                # Rule 2: Zip Bomb Ratio (> 50.0)
                if ratio > MAX_ZIP_BOMB_RATIO:
                    findings.append(
                        CloakedSpanFinding(
                            modality="archive",
                            page_or_paragraph=1,
                            cloak_reason="DECOMPRESSION_BOMB_DETECTED",
                            hidden_text=(
                                f"Decompression bomb ratio {ratio:.1f}:1 exceeds safety threshold 50.0:1 "
                                f"(compressed {compressed_len} bytes -> uncompressed {total_uncompressed} bytes)"
                            ),
                            metadata_key="compression_ratio",
                            risk_category_ids=[8, 19],
                        )
                    )

                # Rule 3: Nested Archives & Path Traversal (ZipSlip)
                for info in infolist:
                    fn_lower = info.filename.lower()
                    if any(fn_lower.endswith(ext) for ext in _ARCHIVE_EXTENSIONS):
                        findings.append(
                            CloakedSpanFinding(
                                modality="archive",
                                page_or_paragraph=1,
                                cloak_reason="SUSPICIOUS_NESTED_ARCHIVE",
                                hidden_text=f"Nested archive file detected inside zip: '{info.filename}'",
                                metadata_key="nested_archive",
                                risk_category_ids=[8],
                            )
                        )
                    if "../" in info.filename or "..\\" in info.filename:
                        findings.append(
                            CloakedSpanFinding(
                                modality="archive",
                                page_or_paragraph=1,
                                cloak_reason="SUSPICIOUS_NESTED_ARCHIVE",
                                hidden_text=f"Directory traversal path in archive member: '{info.filename}'",
                                metadata_key="zip_slip",
                                risk_category_ids=[8],
                            )
                        )

                # Safe summary of contents for Task Continuity
                member_names = [f"- {info.filename} ({info.file_size} bytes)" for info in infolist[:30]]
                summary_text = f"Archive contents safe summary ({len(infolist)} files):\n" + "\n".join(member_names)

        except zipfile.BadZipFile as bz:
            findings.append(
                CloakedSpanFinding(
                    modality="archive",
                    page_or_paragraph=1,
                    cloak_reason="SUSPICIOUS_NESTED_ARCHIVE",
                    hidden_text=f"Corrupted or invalid zip archive payload: {bz}",
                    metadata_key="bad_zip",
                    risk_category_ids=[8],
                )
            )
            summary_text = f"Corrupted or invalid zip archive: {bz}"

        return findings, summary_text

    def scan_spreadsheet_safety(
        self,
        file_bytes: bytes,
        filename: str,
    ) -> Tuple[bytes, List[CloakedSpanFinding], str]:
        """
        Inspect Excel spreadsheets for hidden sheets and malicious formulas
        (=WEBSERVICE, =HYPERLINK, =CMD, =EXEC), unhide data, and scrub formulas.
        """
        findings: List[CloakedSpanFinding] = []
        sanitized_bytes = file_bytes
        visible_text = ""

        fn_lower = filename.lower()
        is_xlsx = fn_lower.endswith((".xlsx", ".xlsm")) or (
            file_bytes.startswith(b"PK\x03\x04") and not fn_lower.endswith(".zip")
        )

        if is_xlsx:
            try:
                wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=False, data_only=False)
                modified = False

                # 1. Unhide hidden or veryHidden worksheets
                for sheet in wb.worksheets:
                    if getattr(sheet, "sheet_state", "visible") in ("hidden", "veryHidden"):
                        findings.append(
                            CloakedSpanFinding(
                                modality="csv_xlsx",
                                page_or_paragraph=1,
                                cloak_reason="HIDDEN_SPREADSHEET_DATA",
                                hidden_text=f"Hidden sheet '{sheet.title}' unhidden for security transparency",
                                metadata_key=f"sheet_{sheet.title}",
                                risk_category_ids=[7],
                            )
                        )
                        sheet.sheet_state = "visible"
                        modified = True

                # 2. Scan and redact malicious formulas
                for sheet in wb.worksheets:
                    cell_count = 0
                    for row in sheet.iter_rows():
                        for cell in row:
                            cell_count += 1
                            if cell_count > 1000:
                                break
                            raw_val = cell.value
                            if raw_val is not None and isinstance(raw_val, str):
                                str_val = raw_val.strip()
                                if _MALICIOUS_FORMULA_RE.search(str_val):
                                    findings.append(
                                        CloakedSpanFinding(
                                            modality="csv_xlsx",
                                            page_or_paragraph=cell.row,
                                            cloak_reason="MALICIOUS_CELL_FORMULA",
                                            hidden_text=(
                                                f"Malicious formula stripped at cell {cell.coordinate}: '{str_val}'"
                                            ),
                                            metadata_key=f"cell_{cell.coordinate}",
                                            risk_category_ids=[1, 7],
                                        )
                                    )
                                    cell.value = "[REDACTED_MALICIOUS_FORMULA]"
                                    modified = True
                        if cell_count > 1000:
                            break

                # Re-save workbook if modified or for consistent clean export
                out_buf = io.BytesIO()
                wb.save(out_buf)
                sanitized_bytes = out_buf.getvalue()

                # Build readable text preview
                preview_parts = []
                for sheet in wb.worksheets:
                    preview_parts.append(f"Sheet: {sheet.title}")
                    for r in list(sheet.iter_rows(values_only=True))[:15]:
                        row_vals = [str(c) for c in r if c is not None]
                        if row_vals:
                            preview_parts.append(" | ".join(row_vals))
                visible_text = "\n".join(preview_parts)

            except Exception as exc:
                logger.debug("openpyxl parse error: %s", exc)
                visible_text = f"Spreadsheet content notice: {exc}"

        elif fn_lower.endswith(".csv"):
            try:
                raw_text = file_bytes.decode("utf-8", errors="replace")
                lines = raw_text.splitlines()
                cleaned_lines = []
                for idx, line in enumerate(lines, start=1):
                    tokens = line.split(",")
                    cleaned_tokens = []
                    for t in tokens:
                        t_str = t.strip()
                        if _MALICIOUS_FORMULA_RE.search(t_str):
                            findings.append(
                                CloakedSpanFinding(
                                    modality="csv_xlsx",
                                    page_or_paragraph=idx,
                                    cloak_reason="MALICIOUS_CELL_FORMULA",
                                    hidden_text=f"Malicious CSV formula at line {idx}: '{t_str}'",
                                    metadata_key=f"csv_line_{idx}",
                                    risk_category_ids=[1, 7],
                                )
                            )
                            cleaned_tokens.append("[REDACTED_MALICIOUS_FORMULA]")
                        else:
                            cleaned_tokens.append(t)
                    cleaned_lines.append(",".join(cleaned_tokens))

                visible_text = "\n".join(cleaned_lines)
                sanitized_bytes = visible_text.encode("utf-8")
            except Exception as exc:
                logger.debug("csv parse error: %s", exc)
                visible_text = f"CSV notice: {exc}"

        return sanitized_bytes, findings, visible_text

    def scan_file_bytes(
        self,
        file_bytes: bytes,
        filename: str = "file.bin",
        modality: Optional[str] = None,
    ) -> DocumentDeCloakResult:
        """
        Main entry point for deep-packet inspection of files, archives, and spreadsheets.
        """
        raw_evidence_sha256 = hashlib.sha256(file_bytes).hexdigest()
        fn_lower = filename.lower()

        # Determine target modality
        if fn_lower.endswith((".zip", ".tar", ".gz", ".7z", ".rar")) or modality == "archive":
            target_modality = "archive"
        elif fn_lower.endswith((".xlsx", ".xlsm", ".csv")) or modality in ("csv_xlsx", "xlsx", "csv"):
            target_modality = "csv_xlsx"
        else:
            target_modality = "file"

        all_findings: List[CloakedSpanFinding] = []
        sanitized_meta: Dict[str, str] = {}
        clean_text = ""
        current_bytes = file_bytes

        # 1. Inspect Magic Bytes and Executable Disguises
        magic_findings = self.scan_magic_bytes(current_bytes, filename)
        all_findings.extend(magic_findings)

        # If disguised executable, immediately return threat result
        if any(f.cloak_reason == "DISGUISED_EXECUTABLE" for f in magic_findings):
            all_rids = [1, 2, 8]
            diff_record = f"ORIGINAL:{raw_evidence_sha256}|BLOCKED_EXECUTABLE"
            redacted_diff_sha256 = hashlib.sha256(diff_record.encode("utf-8")).hexdigest()
            return DocumentDeCloakResult(
                raw_evidence_sha256=raw_evidence_sha256,
                redacted_diff_sha256=redacted_diff_sha256,
                modality=target_modality,  # type: ignore[arg-type]
                visible_clean_text="[BLOCKED: Disguised Executable Binary Intercepted]",
                raw_unfiltered_text="[BINARY_EXECUTABLE_PAYLOAD]",
                cloaked_findings=all_findings,
                metadata_sanitized={"threat": "DISGUISED_EXECUTABLE", "blocked": "true"},
                threat_detected=True,
                risk_category_ids=all_rids,
                summary=f"Hard-blocked disguised executable file: '{filename}'",
            )

        # 2. Archive Safety Check
        if target_modality == "archive" or current_bytes.startswith(b"PK\x03\x04") and not fn_lower.endswith((".xlsx", ".docx")):
            archive_findings, archive_summary = self.scan_archive_safety(current_bytes, filename)
            all_findings.extend(archive_findings)
            clean_text = archive_summary

        # 3. Spreadsheet Safety Check
        if target_modality == "csv_xlsx" or fn_lower.endswith((".xlsx", ".xlsm", ".csv")):
            sanitized_bytes, sheet_findings, sheet_preview = self.scan_spreadsheet_safety(current_bytes, filename)
            all_findings.extend(sheet_findings)
            current_bytes = sanitized_bytes
            clean_text = sheet_preview
            sanitized_meta["sanitized_bytes_b64"] = base64.b64encode(sanitized_bytes).decode("ascii")

        # Aggregate unique risk category IDs
        all_risk_ids: List[int] = []
        for f in all_findings:
            for rid in f.risk_category_ids:
                if rid not in all_risk_ids:
                    all_risk_ids.append(rid)

        threat_detected = len(all_findings) > 0
        diff_record = f"ORIGINAL:{raw_evidence_sha256}|SANITIZED:{clean_text}"
        redacted_diff_sha256 = hashlib.sha256(diff_record.encode("utf-8")).hexdigest()

        reasons = [f.cloak_reason for f in all_findings]
        summary = (
            f"Archive/File scanner detected threats: {reasons}"
            if threat_detected
            else "File/Archive verified clean: valid format and safe contents."
        )

        return DocumentDeCloakResult(
            raw_evidence_sha256=raw_evidence_sha256,
            redacted_diff_sha256=redacted_diff_sha256,
            modality=target_modality,  # type: ignore[arg-type]
            visible_clean_text=clean_text or f"Verified safe file '{filename}'",
            raw_unfiltered_text=f"[BINARY_DATA: {len(file_bytes)} bytes]",
            cloaked_findings=all_findings,
            metadata_sanitized=sanitized_meta,
            threat_detected=threat_detected,
            risk_category_ids=all_risk_ids,
            summary=summary,
        )


# Global singleton instance
archive_file_scanner = ArchiveFileScanner()
